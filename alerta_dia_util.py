# -*- coding: utf-8 -*-
"""Titulo em aberto vencendo em sabado, domingo ou feriado.

Banco nao compensa em dia nao util: o pagamento agendado para sabado, domingo
ou feriado ou nao sai, ou sai no dia seguinte com juros de um dia que ninguem
provisionou. O pedido do e-mail e simples -- ajustar o vencimento para o
PROXIMO DIA UTIL, que ja vai calculado em cada linha.

Le a SE2 direto (nao o painel): quem manda na data e o titulo, nao o boleto.

  em aberto   Saldo > 0 E sem DT Baixa. Titulo pago nao tem o que ajustar.
  a data      "Vencto Real" -- e' a que o financeiro usa para pagar. O
              "Vencimento" contratual pode cair no fim de semana sem problema
              nenhum; o que nao pode e' o REAL. Quando o real vem vazio, cai
              para o contratual (senao o titulo sumiria do alerta calado).
  a janela    7 dias para tras (o que venceu e ficou) e 15 para frente (da
              tempo de ajustar antes de virar corre-corre).

Uso:
    python alerta_dia_util.py                  # manda, no maximo uma vez por dia
    python alerta_dia_util.py --teste          # nao manda: grava a previa
    python alerta_dia_util.py --forcar         # manda de novo no mesmo dia
    python alerta_dia_util.py --dias-atras 7 --dias-frente 15
    python alerta_dia_util.py --base "C:\\caminho\\SE2.xlsx"

Codigo de saida: SEMPRE 0 quando conseguiu ler a base. Falha de e-mail nao
derruba a rodada do painel.
"""

from __future__ import annotations

import argparse
import datetime as dt
import shutil
import sys
import tempfile
from pathlib import Path

import alertas_comuns as comuns

PASTA = Path(__file__).resolve().parent
PREVIA = PASTA / "DADOS" / "alerta_dia_util_previa.html"
CHAVE = "dia_util"

# A SE2 mora fora desta pasta, na biblioteca compartilhada. Derivado, e nao
# escrito na mao: o nome da pasta muda de maquina para maquina.
#   PASTA  ...\LUCAS ABNER ARAUJO\<AUTOMACOES>\ANALISES BOLETOS\PAINEL ANALISE BOLETOS
RAIZ = PASTA.parent.parent.parent
BASES = RAIZ / "BASES GENERICOS"

ABA = "SE2"

# Indices da SE2 exportada (261 colunas). Conferidos pelo NOME no cabecalho na
# hora de ler -- o TOTVS ja inseriu coluna no meio antes, e posicao fixa erraria
# tudo em silencio.
COLUNAS = {
    "filial": "Filial", "prefixo": "Prefixo", "tipo": "Tipo",
    "num": "No. Titulo", "parcela": "Parcela", "fornecedor": "Fornecedor",
    "nome": "Nome Fornece", "emissao": "DT Emissao", "vencimento": "Vencimento",
    "real": "Vencto Real", "valor": "Vlr.Titulo", "baixa": "DT Baixa",
    "saldo": "Saldo", "natureza": "Natureza", "portador": "Portador",
    "uuid": "Campo UUID",
}


def achar_base() -> Path | None:
    """A SE2 mais recente. O nome tem cedilha e til de verdade -- por isso o
    curinga, e nao o nome escrito aqui: acertar o acento depende da pagina de
    codigo com que este arquivo for lido."""
    if not BASES.is_dir():
        return None
    achadas = [p for p in BASES.glob("SE2 - POSI*O DIARIA.xlsx")
               if "backup" not in p.name.lower() and "copia" not in p.name.lower()]
    return max(achadas, key=lambda p: p.stat().st_mtime) if achadas else None


def como_data(valor):
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    if isinstance(valor, str):
        for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
            try:
                return dt.datetime.strptime(valor.strip(), formato).date()
            except ValueError:
                pass
    return None


def como_numero(valor) -> float:
    if isinstance(valor, (int, float)):
        return float(valor)
    if isinstance(valor, str):
        texto = valor.strip().replace(".", "").replace(",", ".")
        try:
            return float(texto)
        except ValueError:
            return 0.0
    return 0.0


def reais(valor: float) -> str:
    return "R$ " + f"{valor:,.2f}".replace(",", "~").replace(".", ",").replace("~", ".")


def ler_se2(caminho: Path, inicio: dt.date, fim: dt.date) -> tuple[list[dict], dict]:
    """Varre a SE2 e devolve o que vence em dia nao util, mais um resumo.

    ⚠ A SE2 costuma estar ABERTA no Excel na hora da rodada, e ai o openpyxl
    leva PermissionError. Copiar para a pasta temporaria e ler a copia resolve
    -- e' leitura, nao muda nada no original.
    """
    import openpyxl

    origem = caminho
    temporario = None
    try:
        arquivo = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    except PermissionError:
        temporario = Path(tempfile.gettempdir()) / f"se2_alerta_{dt.date.today():%Y%m%d}.xlsx"
        shutil.copy2(caminho, temporario)
        arquivo = openpyxl.load_workbook(temporario, read_only=True, data_only=True)
        print(f"   (a SE2 estava aberta; li uma cópia)")

    if ABA not in arquivo.sheetnames:
        arquivo.close()
        raise KeyError(f'a planilha nao tem a aba "{ABA}" (tem: {arquivo.sheetnames})')
    pagina = arquivo[ABA]
    linhas = pagina.iter_rows(values_only=True)
    cabecalho = [str(c).strip() if c is not None else "" for c in next(linhas)]

    onde = {}
    faltando = []
    for chave, rotulo in COLUNAS.items():
        if rotulo in cabecalho:
            onde[chave] = cabecalho.index(rotulo)
        else:
            faltando.append(rotulo)
    # Sem data ou sem saldo o alerta nao tem como existir; o resto e' enfeite.
    for obrigatoria in ("real", "vencimento", "saldo", "num"):
        if obrigatoria not in onde:
            arquivo.close()
            raise KeyError("colunas ausentes na SE2: " + ", ".join(faltando))

    def campo(linha, chave):
        i = onde.get(chave)
        return linha[i] if i is not None and i < len(linha) else None

    achados = []
    resumo = {"linhas": 0, "em_aberto": 0, "na_janela": 0}
    for linha in linhas:
        if campo(linha, "num") in (None, ""):
            continue
        resumo["linhas"] += 1
        if como_numero(campo(linha, "saldo")) <= 0 or como_data(campo(linha, "baixa")):
            continue
        resumo["em_aberto"] += 1
        # o real e' quem manda; o contratual so entra quando o real vem vazio
        data = como_data(campo(linha, "real")) or como_data(campo(linha, "vencimento"))
        if not data or not (inicio <= data <= fim):
            continue
        resumo["na_janela"] += 1
        motivo = comuns.porque_nao_e_util(data)
        if not motivo:
            continue
        contratual = como_data(campo(linha, "vencimento"))
        saldo = como_numero(campo(linha, "saldo"))
        achados.append({
            "uuid": str(campo(linha, "uuid") or "").strip(),
            "filial": str(campo(linha, "filial") or "").strip(),
            "prefixo": str(campo(linha, "prefixo") or "").strip(),
            "tipo": str(campo(linha, "tipo") or "").strip(),
            "titulo": str(campo(linha, "num") or "").strip(),
            "parcela": str(campo(linha, "parcela") or "").strip(),
            "fornecedor": str(campo(linha, "nome") or campo(linha, "fornecedor") or "").strip(),
            "vencimento": f"{data:%d/%m/%Y}",
            "motivo": motivo,
            "ajustar_para": f"{comuns.proximo_dia_util(data):%d/%m/%Y}",
            "contratual": f"{contratual:%d/%m/%Y}" if contratual else "",
            "saldo": reais(saldo),
            "_data": data, "_saldo": saldo,
        })
    arquivo.close()
    if temporario and temporario.exists():
        try:
            temporario.unlink()
        except OSError:
            pass
    resumo["base"] = origem
    resumo["salva_em"] = dt.datetime.fromtimestamp(origem.stat().st_mtime)
    # data primeiro, maior valor antes: o que dói mais aparece no topo do dia
    achados.sort(key=lambda a: (a["_data"], -a["_saldo"]))
    return achados, resumo


COLUNAS_EMAIL = [
    ("Vencimento", "vencimento"),
    ("Cai em", "motivo"),
    ("Ajustar para", "ajustar_para"),
    ("Saldo em aberto", "saldo"),
    ("Filial", "filial"),
    ("Nº título", "titulo"),
    ("Parc.", "parcela"),
    ("Tipo", "tipo"),
    ("Fornecedor", "fornecedor"),
    ("Código (UUID)", "uuid"),
]


def montar_html(achados: list[dict], resumo: dict, inicio: dt.date, fim: dt.date) -> str:
    quantos = len(achados)
    total = sum(a["_saldo"] for a in achados)
    frase = ("1 título em aberto está vencendo em dia não útil"
             if quantos == 1 else
             f"{quantos} títulos em aberto estão vencendo em dia não útil")
    return f"""<div style="{comuns.FONTE}font-size:11pt;color:#000000;">
  <p style="margin:0 0 12px 0;">Bom dia, Graziela,</p>

  <p style="margin:0 0 12px 0;">{frase}, somando
  <b>{reais(total)}</b>.</p>

  <p style="margin:0 0 14px 0;">O banco não compensa em sábado, domingo ou
  feriado. <b>Favor ajustar o vencimento para o próximo dia útil</b> — a data
  sugerida já vai na coluna <i>Ajustar para</i>.</p>

  {comuns.tabela(COLUNAS_EMAIL, achados,
                 direita=("saldo",), destaque=("vencimento", "ajustar_para", "uuid"))}

  {comuns.rodape(
      f"Janela: vencimentos de {inicio:%d/%m/%Y} a {fim:%d/%m/%Y} "
      f"(7 dias para trás e 15 para frente).",
      f"Base SE2 salva em {resumo['salva_em']:%d/%m/%Y às %H:%M} — "
      f"{resumo['em_aberto']} títulos em aberto.")}
</div>"""


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--para", default="")
    parser.add_argument("--teste", action="store_true",
                        help="nao envia: grava a previa e mostra o resumo")
    parser.add_argument("--forcar", action="store_true",
                        help="envia mesmo se ja mandou hoje")
    parser.add_argument("--dias-atras", type=int, default=7)
    parser.add_argument("--dias-frente", type=int, default=15)
    parser.add_argument("--base", type=Path, default=None)
    args = parser.parse_args()

    base = args.base or achar_base()
    if not base or not base.exists():
        print(f"   AVISO: nao achei a SE2 em {BASES} -- alerta de dia util nao rodou.")
        return 0

    hoje = dt.date.today()
    inicio = hoje - dt.timedelta(days=args.dias_atras)
    fim = hoje + dt.timedelta(days=args.dias_frente)

    try:
        achados, resumo = ler_se2(base, inicio, fim)
    except Exception as erro:  # noqa: BLE001
        print(f"   AVISO: nao consegui ler a SE2 ({erro}). A rodada segue.")
        return 0

    print(f"   SE2: {resumo['linhas']} títulos | {resumo['em_aberto']} em aberto | "
          f"{resumo['na_janela']} vencendo entre {inicio:%d/%m} e {fim:%d/%m}")

    if not achados:
        print("   Nenhum título em aberto vencendo em sábado, domingo ou feriado.")
        return 0

    total = sum(a["_saldo"] for a in achados)
    print(f"   {len(achados)} título(s) em dia NÃO ÚTIL ({reais(total)}):")
    for a in achados:
        print(f"     {a['vencimento']} {a['motivo']:<12} -> {a['ajustar_para']} | "
              f"{a['filial']} {a['titulo']}/{a['parcela'] or '-'} | "
              f"{a['saldo']:>16} | {a['fornecedor'][:30]}")

    corpo = montar_html(achados, resumo, inicio, fim)
    assunto = (f"[Painel Boletos] {len(achados)} título"
               f"{'s' if len(achados) > 1 else ''} vencendo em dia não útil - "
               f"ajustar para o próximo dia útil - {hoje:%d/%m/%Y}")

    if args.teste:
        PREVIA.parent.mkdir(parents=True, exist_ok=True)
        PREVIA.write_text(corpo, encoding="utf-8")
        print(f"\n   MODO TESTE: nada foi enviado.\n   Assunto: {assunto}")
        print(f"   Prévia:  {PREVIA}")
        return 0

    destino = args.para or comuns.destinatarios(CHAVE)
    if not destino:
        print("   AVISO: nao achei para quem mandar (.alerta_adiantamento_para).")
        return 0

    if not args.forcar and comuns.ja_enviado_hoje(CHAVE):
        print(f"   Já enviado hoje ({comuns.quando_enviou(CHAVE)}) -- não mando de novo.")
        print("   Para mandar assim mesmo: --forcar")
        return 0

    try:
        comuns.enviar(assunto, corpo, destino)
        comuns.marcar_enviado(CHAVE, len(achados))
        print(f"   E-mail enviado para {destino}.")
    except Exception as erro:  # noqa: BLE001
        print(f"   AVISO: o e-mail NAO foi enviado ({erro}). O painel segue.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
