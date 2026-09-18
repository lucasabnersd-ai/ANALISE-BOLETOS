# -*- coding: utf-8 -*-
"""Titulo com Tipo Pgto = BOLETO S/C que TEM linha digitavel e codigo de barras.

A contradicao e o alerta: o campo "Tipo Pgto" diz BOLETO S/C, mas o titulo esta
com "Linha Dig." E "Cod.Barras" preenchidos na SE2. Ou a classificacao esta
errada, ou os dados do boleto entraram num titulo que nao deveria te-los -- e
quem paga precisa saber qual das duas antes de pagar. O e-mail so apresenta o
fato e pede a verificacao; nao decide qual lado esta errado.

  o filtro    Tipo Pgto (coluna IK) == "BOLETO S/C", comparado em MAIUSCULA e
              sem espaco nas pontas -- o TOTVS devolve o mesmo tipo com caixa
              e espaco variando, e comparar cru deixaria titulo de fora calado.
  os dados    "Linha Dig." (DU) e "Cod.Barras" (DT), os DOIS preenchidos. Um
              so nao e contradicao: e cadastro pela metade, outro assunto.
  em aberto   Saldo > 0 E sem DT Baixa. Titulo ja pago nao tem o que verificar
              -- em 18/09/2026 eram 65 casos na base inteira e so 18 em aberto.
  a data      "Vencto Real", caindo para o "Vencimento" contratual quando o
              real vem vazio (a mesma regra do alerta_dia_util: o real e' o que
              o financeiro usa para pagar).
  a janela    32 dias para tras e 7 para frente -- pedido dele em 18/09/2026.
              Para tras pega o que venceu e ficou parado sem ninguem resolver;
              para frente, o que ainda da tempo de corrigir antes de pagar.

Vai para o Rafael e a Graziela, com a Gabriella em copia
(`.alerta_boleto_sc_para` e `.alerta_boleto_sc_copia`, os dois fora do repo).

Uso:
    python alerta_boleto_sc.py                 # manda, no maximo uma vez por dia
    python alerta_boleto_sc.py --teste         # nao manda: grava a previa
    python alerta_boleto_sc.py --forcar        # manda de novo no mesmo dia
    python alerta_boleto_sc.py --dias-atras 32 --dias-frente 7
    python alerta_boleto_sc.py --base "C:\\caminho\\SE2.xlsx"

Codigo de saida: SEMPRE 0 quando conseguiu ler a base. Falha de e-mail nao
derruba a rodada do painel.
"""

from __future__ import annotations

import argparse
import datetime as dt
import shutil
import sys
import tempfile
from pathlib import Path

import alertas_comuns as comuns
import caminhos

PASTA = Path(__file__).resolve().parent
PREVIA = PASTA / "DADOS" / "alerta_boleto_sc_previa.html"
CHAVE = "boleto_sc"

RAIZ = caminhos.raiz_lucas()
BASES = caminhos.bases_genericos()

ABA = "SE2"

# O tipo procurado. Comparado normalizado (upper + strip) em vez de cru.
TIPO_ALVO = "BOLETO S/C"

# O painel da SE2 no ar. O numero do titulo na tabela vira link para ca -- pedido
# dele em 18/09/2026 -- para quem recebe abrir a SE2 e conferir o titulo sem ter
# de procurar o endereco. Endereco da RAIZ de proposito: o painel nao le filtro
# pela URL (conferido no index.html em 18/09/2026), entao link por titulo cairia
# na mesma tela e so daria a impressao falsa de ja vir filtrado.
PAINEL_SE2 = "https://se2-lucas.vercel.app/"

# Conferidas pelo NOME no cabecalho, nunca por posicao: o TOTVS ja inseriu
# coluna no meio da exportacao antes, e indice fixo erraria tudo em silencio.
COLUNAS = {
    "filial": "Filial", "prefixo": "Prefixo", "tipo": "Tipo",
    "num": "No. Titulo", "parcela": "Parcela", "fornecedor": "Fornecedor",
    "nome": "Nome Fornece", "razao": "Razão Social",
    "vencimento": "Vencimento", "real": "Vencto Real",
    "valor": "Vlr.Titulo", "baixa": "DT Baixa", "saldo": "Saldo",
    "tipo_pgto": "Tipo Pgto", "linha_dig": "Linha Dig.",
    "cod_barras": "Cod.Barras", "uuid": "Campo UUID",
}


def achar_base() -> Path | None:
    """A SE2 mais recente. O nome tem cedilha e til de verdade -- por isso o
    curinga, e nao o nome escrito aqui."""
    if not BASES.is_dir():
        return None
    achadas = [p for p in BASES.glob("SE2 - POSI*O DIARIA.xlsx")
               if "backup" not in p.name.lower() and "copia" not in p.name.lower()]
    return max(achadas, key=lambda p: p.stat().st_mtime) if achadas else None


def como_data(valor):
    if isinstance(valor, dt.datetime):
        return valor.date()
    if isinstance(valor, dt.date):
        return valor
    if isinstance(valor, str):
        for formato in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
            try:
                return dt.datetime.strptime(valor.strip(), formato).date()
            except ValueError:
                pass
    return None


def como_numero(valor) -> float:
    if isinstance(valor, (int, float)):
        return float(valor)
    if isinstance(valor, str):
        texto = valor.strip().replace(".", "").replace(",", ".")
        try:
            return float(texto)
        except ValueError:
            return 0.0
    return 0.0


def reais(valor: float) -> str:
    return "R$ " + f"{valor:,.2f}".replace(",", "~").replace(".", ",").replace("~", ".")


def texto(valor) -> str:
    """Texto util da celula. O TOTVS devolve o campo vazio de varios jeitos --
    None, a string "None", data em branco "  /  /    " -- e todos significam
    a mesma coisa: nao tem nada ali. Tratar so o None deixaria passar
    "None" como se fosse linha digitavel preenchida."""
    if valor is None:
        return ""
    limpo = str(valor).strip()
    if limpo.lower() in ("none", "nan", "null"):
        return ""
    # so barra, espaco e dois-pontos = data/hora em branco do TOTVS
    if limpo and set(limpo) <= set(" /:-"):
        return ""
    return limpo


def ler_se2(caminho: Path, inicio: dt.date, fim: dt.date) -> tuple[list[dict], dict]:
    """Varre a SE2 e devolve os BOLETO S/C contraditorios, mais um resumo.

    ⚠ A SE2 costuma estar ABERTA no Excel na hora da rodada, e ai o openpyxl
    leva PermissionError. Copiar para a pasta temporaria e ler a copia resolve
    -- e' leitura, nao muda nada no original.
    """
    import openpyxl

    origem = caminho
    temporario = None
    try:
        arquivo = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    except PermissionError:
        temporario = Path(tempfile.gettempdir()) / f"se2_sc_{dt.date.today():%Y%m%d}.xlsx"
        shutil.copy2(caminho, temporario)
        arquivo = openpyxl.load_workbook(temporario, read_only=True, data_only=True)
        print("   (a SE2 estava aberta; li uma cópia)")

    if ABA not in arquivo.sheetnames:
        arquivo.close()
        raise KeyError(f'a planilha nao tem a aba "{ABA}" (tem: {arquivo.sheetnames})')
    pagina = arquivo[ABA]
    linhas = pagina.iter_rows(values_only=True)
    cabecalho = [str(c).strip() if c is not None else "" for c in next(linhas)]

    onde = {}
    faltando = []
    for chave, rotulo in COLUNAS.items():
        if rotulo in cabecalho:
            onde[chave] = cabecalho.index(rotulo)
        else:
            faltando.append(rotulo)
    # Sem estas o alerta nao existe. Faltando uma, ERRA ALTO em vez de mandar
    # e-mail vazio: "nenhum titulo" e "nao consegui olhar" sao coisas
    # diferentes, e so a segunda precisa de conserto.
    for obrigatoria in ("tipo_pgto", "linha_dig", "cod_barras", "saldo", "num"):
        if obrigatoria not in onde:
            arquivo.close()
            raise KeyError("colunas ausentes na SE2: " + ", ".join(faltando))

    def campo(linha, chave):
        i = onde.get(chave)
        return linha[i] if i is not None and i < len(linha) else None

    achados = []
    resumo = {"linhas": 0, "em_aberto": 0, "do_tipo": 0, "na_janela": 0}
    for linha in linhas:
        if campo(linha, "num") in (None, ""):
            continue
        resumo["linhas"] += 1
        if texto(campo(linha, "tipo_pgto")).upper() != TIPO_ALVO:
            continue
        resumo["do_tipo"] += 1
        if como_numero(campo(linha, "saldo")) <= 0 or como_data(campo(linha, "baixa")):
            continue
        resumo["em_aberto"] += 1
        data = como_data(campo(linha, "real")) or como_data(campo(linha, "vencimento"))
        if not data or not (inicio <= data <= fim):
            continue
        resumo["na_janela"] += 1

        linha_dig = texto(campo(linha, "linha_dig"))
        cod_barras = texto(campo(linha, "cod_barras"))
        if not (linha_dig and cod_barras):
            continue

        saldo = como_numero(campo(linha, "saldo"))
        achados.append({
            "uuid": texto(campo(linha, "uuid")),
            "filial": texto(campo(linha, "filial")),
            "prefixo": texto(campo(linha, "prefixo")),
            "tipo": texto(campo(linha, "tipo")),
            "titulo": texto(campo(linha, "num")),
            "parcela": texto(campo(linha, "parcela")),
            "fornecedor": texto(campo(linha, "razao")) or texto(campo(linha, "nome")),
            "vencimento": f"{data:%d/%m/%Y}",
            "tipo_pgto": texto(campo(linha, "tipo_pgto")),
            "linha_dig": linha_dig,
            "cod_barras": cod_barras,
            "saldo": reais(saldo),
            "_data": data, "_saldo": saldo,
        })
    arquivo.close()
    if temporario and temporario.exists():
        try:
            temporario.unlink()
        except OSError:
            pass
    resumo["base"] = origem
    resumo["salva_em"] = dt.datetime.fromtimestamp(origem.stat().st_mtime)
    # o mais antigo primeiro: o que venceu e ficou parado e o que mais incomoda
    achados.sort(key=lambda a: (a["_data"], -a["_saldo"]))
    return achados, resumo


COLUNAS_EMAIL = [
    ("Vencimento", "vencimento"),
    ("Filial", "filial"),
    ("Prefixo", "prefixo"),
    ("Tipo", "tipo"),
    ("Nº título", "titulo"),
    ("Parc.", "parcela"),
    ("Fornecedor", "fornecedor"),
    ("Saldo em aberto", "saldo"),
    ("Código (UUID)", "uuid"),
]
# ⚠ "Linha digitável" e "Tipo pgto" ficam FORA da tabela de proposito, apesar de
# serem o criterio. Conferido renderizado em 18/09/2026: a linha digitavel sao 47
# digitos numa palavra so, que nao quebra -- ela sozinha empurrava a tabela para
# ~1.100px e jogava a coluna do UUID para fora da tela, justo a que ele pediu. O
# "Tipo pgto" e constante (todo mundo na tabela e BOLETO S/C) e ja esta na frase
# acima. Os dois continuam sendo lidos e conferidos; so nao viram coluna.


def montar_html(achados: list[dict], resumo: dict, inicio: dt.date, fim: dt.date) -> str:
    quantos = len(achados)
    total = sum(a["_saldo"] for a in achados)
    frase = ("1 título em aberto está classificado como"
             if quantos == 1 else
             f"{quantos} títulos em aberto estão classificados como")
    plural = "esse título" if quantos == 1 else "esses títulos"
    return f"""<div style="{comuns.FONTE}font-size:11pt;color:#000000;">
  <p style="margin:0 0 12px 0;">Bom dia,</p>

  <p style="margin:0 0 12px 0;">{frase} <b>{TIPO_ALVO}</b> na SE2, mas
  <b>com linha digitável e código de barras preenchidos</b> — somando
  <b>{reais(total)}</b> em aberto.</p>

  <p style="margin:0 0 14px 0;">São duas informações que se contradizem: ou o
  <i>Tipo pgto</i> está errado, ou os dados do boleto entraram num título que
  não deveria tê-los. <b>Favor verificar {plural} antes do pagamento</b> e
  acertar o lado que estiver incorreto. O <b>nº do título</b> na tabela abre o
  <a href="{PAINEL_SE2}" style="color:#1F3864;">painel da SE2</a>, onde dá para
  procurar pelo código (UUID) da última coluna.</p>

  {comuns.tabela(COLUNAS_EMAIL, achados,
                 direita=("saldo",),
                 destaque=("vencimento", "prefixo", "tipo", "uuid"),
                 links={"titulo": PAINEL_SE2})}

  {comuns.rodape(
      f"Janela: vencimentos de {inicio:%d/%m/%Y} a {fim:%d/%m/%Y} "
      f"(32 dias para trás e 7 para frente), somente títulos sem baixa.",
      f"Base SE2 salva em {resumo['salva_em']:%d/%m/%Y às %H:%M} — "
      f"{resumo['do_tipo']} títulos {TIPO_ALVO} na base, "
      f"{resumo['em_aberto']} em aberto.")}
</div>"""


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--para", default="")
    parser.add_argument("--copia", default="")
    parser.add_argument("--teste", action="store_true",
                        help="nao envia: grava a previa e mostra o resumo")
    parser.add_argument("--forcar", action="store_true",
                        help="envia mesmo se ja mandou hoje")
    parser.add_argument("--dias-atras", type=int, default=32)
    parser.add_argument("--dias-frente", type=int, default=7)
    parser.add_argument("--base", type=Path, default=None)
    args = parser.parse_args()

    base = args.base or achar_base()
    if not base or not base.exists():
        print(f"   AVISO: nao achei a SE2 em {BASES} -- alerta BOLETO S/C nao rodou.")
        return 0

    hoje = dt.date.today()
    inicio = hoje - dt.timedelta(days=args.dias_atras)
    fim = hoje + dt.timedelta(days=args.dias_frente)

    try:
        achados, resumo = ler_se2(base, inicio, fim)
    except Exception as erro:  # noqa: BLE001
        print(f"   AVISO: nao consegui ler a SE2 ({erro}). A rodada segue.")
        return 0

    print(f"   SE2: {resumo['linhas']} títulos | {resumo['do_tipo']} {TIPO_ALVO} | "
          f"{resumo['em_aberto']} em aberto | {resumo['na_janela']} vencendo entre "
          f"{inicio:%d/%m} e {fim:%d/%m}")

    if not achados:
        print(f"   Nenhum título {TIPO_ALVO} em aberto com linha digitável e "
              f"código de barras preenchidos.")
        return 0

    total = sum(a["_saldo"] for a in achados)
    print(f"   {len(achados)} título(s) {TIPO_ALVO} COM dados de boleto ({reais(total)}):")
    for a in achados:
        print(f"     {a['vencimento']} | {a['filial']} | pref {a['prefixo'] or '-':<4} | "
              f"tipo {a['tipo']:<4} | {a['titulo']}/{a['parcela'] or '-'} | "
              f"{a['saldo']:>15} | {a['fornecedor'][:30]} | {a['uuid']}")

    corpo = montar_html(achados, resumo, inicio, fim)
    assunto = (f"[Painel Boletos] {len(achados)} título"
               f"{'s' if len(achados) > 1 else ''} {TIPO_ALVO} com linha digitável e "
               f"código de barras - verificar - {hoje:%d/%m/%Y}")

    if args.teste:
        PREVIA.parent.mkdir(parents=True, exist_ok=True)
        PREVIA.write_text(corpo, encoding="utf-8")
        print(f"\n   MODO TESTE: nada foi enviado.\n   Assunto: {assunto}")
        print(f"   Prévia:  {PREVIA}")
        return 0

    destino = args.para or comuns.destinatarios(CHAVE)
    if not destino:
        print("   AVISO: nao achei para quem mandar (.alerta_boleto_sc_para).")
        return 0
    copia = args.copia or comuns.copias(CHAVE)

    if not args.forcar and comuns.ja_enviado_hoje(CHAVE):
        print(f"   Já enviado hoje ({comuns.quando_enviou(CHAVE)}) -- não mando de novo.")
        print("   Para mandar assim mesmo: --forcar")
        return 0

    try:
        comuns.enviar(assunto, corpo, destino, copia=copia)
        comuns.marcar_enviado(CHAVE, len(achados))
        print(f"   E-mail enviado para {destino}"
              + (f" (cópia: {copia})" if copia else "") + ".")
    except Exception as erro:  # noqa: BLE001
        print(f"   AVISO: o e-mail NAO foi enviado ({erro}). O painel segue.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
