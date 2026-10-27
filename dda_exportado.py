# -*- coding: utf-8 -*-
"""dda_exportado.py - os Excel do DDA exportados pelo banco viram fonte do painel.

Pedido dele em 27/10/2026: "SALVEI AS BASES EXCEL DO BOLETOS VENCIDOS E A VENCER
[em PAINEIS\\ANALISE BOLETOS\\BOLETOS VENCIDOS A VENCER] USE ELES PARA ALIMENTAR AS
ABAS QUE VINHAM DESATUALIZADAS DE BOLETOS EM ABERTO NO DDA E BOLETOS NAO
ASSOCIADOS A VENCER".

Ate aqui a aba "Boletos em Aberto DDA" lia o COPIAR E COLAR BOLETOS PENDENTES.xlsx,
colado a mao -- parado desde 27/08. Os arquivos novos sao a exportacao do proprio
DDA (DDA_ConsultaBoletosRecebidos_DDMMAAAA_HHMMSS.xlsx): 10 linhas de topo (conta,
empresa, data/hora), o cabecalho na linha que tem "Beneficiário", e o CODIGO DE
BARRAS na ultima coluna -- o que o copiar e colar nunca teve.

Duas saidas:
  linhas_brutas()        -> cabecalho + linhas no formato que o pendentes_itau.py
                            ja le (aba "Boletos em Aberto DDA");
  gerar_base_associador() -> DADOS\\DDA_EXPORTADO_PARA_ASSOCIADOR.xlsx, com as
                            colunas da BOLETOS ITAU.xlsx, que o associador le como
                            mais uma fonte (aba "Boletos Nao Associados a Vencer").

Regras:
  - vale SO a exportacao do DIA MAIS RECENTE (a data do nome do arquivo): cada dia
    e uma foto do DDA, e misturar dias ressuscitaria boleto ja pago;
  - os varios arquivos do dia (uma tela por empresa/filtro) sao somados e a
    repeticao sai pelo codigo de barras (sem codigo: CNPJ+doc+venc+valor).

Uso: python dda_exportado.py   (gera a base do associador e diz quantos boletos leu)
"""
from __future__ import annotations

import datetime as dt
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

AQUI = Path(__file__).resolve().parent
PASTA = AQUI / "BOLETOS VENCIDOS A VENCER"
SAIDA_ASSOCIADOR = AQUI / "DADOS" / "DDA_EXPORTADO_PARA_ASSOCIADOR.xlsx"

CAMPOS = ["Pagador/Agregado", "Beneficiário", "CPF/CNPJ", "Venc.", "Nº Doc.",
          "Até Venc.", "A Pagar", "Tipo de Boleto", "Banco", "Observações", "Código de barras"]
_NOME = re.compile(r"DDA_.*_(\d{2})(\d{2})(\d{4})_(\d{6})\.xlsx$", re.I)


def _rot(v) -> str:
    import unicodedata
    s = unicodedata.normalize("NFKD", str(v or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def arquivos_do_dia() -> list[Path]:
    """Os arquivos da exportacao MAIS RECENTE (pela data do nome)."""
    if not PASTA.is_dir():
        return []
    datados = []
    for p in PASTA.glob("*.xlsx"):
        if p.name.startswith("~$"):
            continue
        m = _NOME.search(p.name)
        if m:
            datados.append((dt.date(int(m[3]), int(m[2]), int(m[1])), p))
    if not datados:
        return []
    ultimo = max(d for d, _ in datados)
    return sorted(p for d, p in datados if d == ultimo)


def mais_nova() -> float:
    arqs = arquivos_do_dia()
    return max(p.stat().st_mtime for p in arqs) if arqs else 0.0


def _ler_arquivo(p: Path) -> list[dict]:
    wb = load_workbook(p, read_only=True, data_only=True)
    try:
        linhas = list(wb.worksheets[0].iter_rows(values_only=True))
    finally:
        wb.close()
    ini = next((i for i, r in enumerate(linhas)
                if r and any(_rot(c) == "beneficiario" for c in r)), None)
    if ini is None:
        raise RuntimeError(f"{p.name}: nao achei a linha de cabecalho (coluna 'Beneficiário')")
    cab = {_rot(c): i for i, c in enumerate(linhas[ini]) if c}
    falta = [c for c in CAMPOS if _rot(c) not in cab]
    if falta:
        raise RuntimeError(f"{p.name}: faltam as colunas {falta}")
    saida = []
    for r in linhas[ini + 1:]:
        if not r or all(v in (None, "") for v in r):
            continue
        reg = {c: r[cab[_rot(c)]] if cab[_rot(c)] < len(r) else None for c in CAMPOS}
        if not reg["Beneficiário"] and reg["Até Venc."] in (None, ""):
            continue
        reg["_arquivo"] = p.name
        saida.append(reg)
    return saida


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.-]", "", str(v or ""))
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _data(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    try:
        return dt.datetime.strptime(str(v).strip(), "%d/%m/%Y").date()
    except ValueError:
        return None


def boletos() -> list[dict]:
    """Os boletos do dia mais recente, sem repeticao."""
    vistos, saida = set(), []
    for p in arquivos_do_dia():
        for reg in _ler_arquivo(p):
            barras = re.sub(r"\D", "", str(reg["Código de barras"] or ""))
            chave = barras or "|".join(str(reg[c]) for c in ("CPF/CNPJ", "Nº Doc.", "Venc.", "Até Venc."))
            if chave in vistos:
                continue
            vistos.add(chave)
            saida.append(reg)
    return saida


def _moeda_br(v) -> str:
    n = _num(v)
    if n is None:
        return ""
    return "R$ " + f"{n:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def linhas_brutas() -> list[tuple]:
    """Cabecalho + linhas no layout do DDA que o pendentes_itau.py ja le.
    Valores em texto "R$ 1.440,00", igual ao copiar e colar: a CHAVE do boleto
    usa os digitos do valor, e numero cru (1440.0) mudaria a chave de todos."""
    cab = CAMPOS[:10]
    saida = [tuple(cab)]
    for reg in boletos():
        linha = []
        for c in cab:
            v = reg[c]
            if c in ("Até Venc.", "A Pagar"):
                v = _moeda_br(v)
            elif c == "Venc.":
                v = _data(v) or v
            linha.append(v)
        saida.append(tuple(linha))
    return saida


def _mod10(numero: str) -> str:
    soma, peso = 0, 2
    for d in reversed(numero):
        p = int(d) * peso
        soma += p // 10 + p % 10
        peso = 1 if peso == 2 else 2
    return str((10 - soma % 10) % 10)


def linha_digitavel(barras: str) -> str:
    """Linha digitavel (47) a partir do codigo de barras (44) de boleto bancario.
    Convenio/concessionaria (comeca com 8) fica em branco: tem outra regra."""
    b = re.sub(r"\D", "", str(barras or ""))
    if len(b) != 44 or b.startswith("8"):
        return ""
    livre = b[19:44]
    c1 = b[0:4] + livre[0:5]
    c2 = livre[5:15]
    c3 = livre[15:25]
    c1 += _mod10(c1); c2 += _mod10(c2); c3 += _mod10(c3)
    return (f"{c1[:5]}.{c1[5:]} {c2[:5]}.{c2[5:]} {c3[:5]}.{c3[5:]} "
            f"{b[4]} {b[5:19]}")


def gerar_base_associador(destino: Path = SAIDA_ASSOCIADOR) -> int:
    """Grava a base no layout da BOLETOS ITAU.xlsx para o associador ler."""
    regs = boletos()
    wb = Workbook()
    ws = wb.active
    ws.title = "Boletos"
    ws.append(["Arquivo", "Empresa Pagadora", "Favorecido/Fornecedor", "CNPJ",
               "Nº Nota/Doc", "Data Vencimento", "Valor (R$)",
               "Código de Barras (44 dígitos)", "Linha Digitável"])
    for r in regs:
        barras = re.sub(r"\D", "", str(r["Código de barras"] or ""))
        ws.append([f"DDA EXPORTADO {r['_arquivo']}", r["Pagador/Agregado"], r["Beneficiário"],
                   re.sub(r"\D", "", str(r["CPF/CNPJ"] or "")), r["Nº Doc."],
                   _data(r["Venc."]), _num(r["Até Venc."]), barras, linha_digitavel(barras)])
    destino.parent.mkdir(parents=True, exist_ok=True)
    tmp = destino.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(destino)
    return len(regs)


if __name__ == "__main__":
    arqs = arquivos_do_dia()
    if not arqs:
        print("DDA exportado: nenhum arquivo em", PASTA)
        sys.exit(0)
    n = gerar_base_associador()
    print(f"DDA exportado: {len(arqs)} arquivo(s) de {arqs[0].name.split('_')[2]} -> "
          f"{n} boletos sem repeticao -> {SAIDA_ASSOCIADOR.name}")
